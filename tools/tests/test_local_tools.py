"""Exercise the public dispatcher with fictional inputs and new local files."""
import csv
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
ENTRY = ROOT / "tools/run_skill.py"


class LocalToolsTest(unittest.TestCase):
    def run_tool(self, tool, kind, source, output, dry_run=False):
        command = [sys.executable, "-B", str(ENTRY), tool, "--type", kind,
                   "--input", str(source), "--output", str(output)]
        if dry_run:
            command.append("--dry-run")
        result = subprocess.run(command, cwd=source.parent, capture_output=True,
                                text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        return json.loads(result.stdout)

    def test_card_dispatcher_preserves_input_and_generates_both_types(self):
        with tempfile.TemporaryDirectory(prefix="partner cards ") as tmp:
            directory = Path(tmp)
            for kind, columns in [("mianyong", 27), ("feizhang", 29)]:
                with self.subTest(kind=kind):
                    record = {"title": "虚构活动", "city_ids": "北京市", "price": "0",
                              "sale_start_time": "2026-10-20 00:00:00", "sale_end_time": "2026-10-20 23:59:59",
                              "effect_start_time": "2026-10-20 00:00:00", "effect_end_time": "2026-10-21 23:59:59",
                              "begin_time_1": "07:00:00", "end_time_1": "08:30:00"}
                    if kind == "feizhang":
                        record.update(reward_percent=5, max_amount_of_order="1.00")
                    source, output = directory / (kind + ".json"), directory / (kind + ".xlsx")
                    source.write_text(json.dumps([record], ensure_ascii=False), encoding="utf-8")
                    original = source.read_bytes()
                    before = set(directory.iterdir())
                    self.run_tool("kaquan", kind, source, output, dry_run=True)
                    self.assertEqual(set(directory.iterdir()), before)
                    result = self.run_tool("kaquan", kind, source, output)
                    self.assertEqual(result["readback"], "passed")
                    self.assertEqual(source.read_bytes(), original)
                    workbook = load_workbook(output)
                    self.assertEqual(workbook.sheetnames, ["活动信息", "重要说明"])
                    sheet = workbook["活动信息"]
                    self.assertEqual((sheet.max_row, sheet.max_column), (3, columns))
                    fields = [cell.value for cell in sheet[2]]
                    self.assertEqual(sheet.cell(3, fields.index("price") + 1).value, "0.00")
                    workbook.close()

    def test_adjustment_dispatcher_generates_credit_and_debit_without_changing_id(self):
        with tempfile.TemporaryDirectory(prefix="partner adjustment ") as tmp:
            directory = Path(tmp)
            for kind, description in [("credit", "补款"), ("debit", "扣款")]:
                with self.subTest(kind=kind):
                    source, output = directory / (kind + ".json"), directory / (kind + ".csv")
                    source.write_text(json.dumps([{"driver_id": "00123456789012345678", "driver_name": "示例司机", "amount_yuan": "25"}]), encoding="utf-8")
                    original = source.read_bytes()
                    before = set(directory.iterdir())
                    self.run_tool("diaozhang", kind, source, output, dry_run=True)
                    self.assertEqual(set(directory.iterdir()), before)
                    result = self.run_tool("diaozhang", kind, source, output)
                    self.assertEqual(result["readback"], "passed")
                    self.assertEqual(source.read_bytes(), original)
                    self.assertTrue(output.read_bytes().startswith(b"\xef\xbb\xbf"))
                    rows = list(csv.DictReader(io.StringIO(output.read_text(encoding="utf-8-sig"))))
                    self.assertEqual(len(rows), 1)
                    self.assertEqual(rows[0]["driver_id"], "00123456789012345678")
                    self.assertEqual(rows[0]["amount_yuan"], "25.00")
                    self.assertEqual(rows[0]["desc"], description)


if __name__ == "__main__":
    unittest.main()
