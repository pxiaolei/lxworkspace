#!/usr/bin/env python3
"""Create a new card-import workbook from a completely validated batch."""

import argparse
import json
import re
from collections import Counter
from copy import copy
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from zipfile import BadZipFile

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


ROOT = Path(__file__).resolve().parents[1]
PREFIX = ["id", "title", "city_ids", "sale_start_time", "sale_end_time", "effect_start_time", "effect_end_time", "tag"]
PERIODS = [field for index in range(1, 6) for field in (f"begin_time_{index}", f"end_time_{index}")]
SUFFIX = ["join_dri_tag", "forbid_dri_tag"] + PERIODS + ["order_category", "attention"]
FIELDS = {
    "mianyong": PREFIX + ["stock", "get_type", "price", "card_rule", "order_limit"] + SUFFIX,
    "feizhang": PREFIX + ["get_type", "stock", "price", "card_rule", "order_limit", "reward_percent", "max_amount_of_order"] + SUFFIX,
}
RULES = {
    "mianyong": ("单单免佣", "每天前x单免佣", "累计前x单免佣"),
    "feizhang": ("单单飞涨", "每天前x单飞涨", "累计前x单飞涨"),
}
REQUIRED = ["title", "city_ids", "price", "sale_start_time", "sale_end_time", "effect_start_time", "effect_end_time"]
CATEGORIES = {"普通单", "特赚单", "快车单", "增量单", "稳收单"}
DEFAULT_FIELDS = {"tag", "get_type", "stock", "card_rule", "order_limit", "order_category", "join_dri_tag", "forbid_dri_tag", "attention"}


def load_assets(kind):
    if load_workbook is None:
        raise ValueError("缺少 openpyxl，请在任务虚拟环境准备 requirements.txt 中的依赖；本入口不会安装依赖")
    wb = load_workbook(ROOT / "assets" / (kind + "_template.xlsx"), data_only=False)
    ws = wb["活动信息"]
    if wb.sheetnames != ["活动信息", "重要说明"] or ws.max_row != 3 or [cell.value for cell in ws[2]] != FIELDS[kind]:
        wb.close()
        raise ValueError("原模板的工作表、示例行或英文字段与已登记契约不一致")
    payload = json.loads((ROOT / "defaults.json").read_text(encoding="utf-8"))
    defaults = payload.get(kind) if isinstance(payload, dict) else None
    if not isinstance(defaults, dict) or set(defaults) != DEFAULT_FIELDS:
        wb.close()
        raise ValueError("defaults.json 必须包含已确认的九个默认字段，不能代填用户必填项")
    defaults = dict(defaults)
    if defaults["attention"] is None:
        defaults["attention"] = ws.cell(3, FIELDS[kind].index("attention") + 1).value
    return wb, defaults


def text(value, field, allow_empty=False):
    if not isinstance(value, str) or (not allow_empty and not value.strip()):
        raise ValueError(f"{field} 必须是用户提供的非空文本" if not allow_empty else f"{field} 必须是文本")
    return value


def integer(value, field, minimum=None, maximum=None):
    if isinstance(value, Decimal) and value.is_finite() and value == value.to_integral_value():
        value = int(value)
    if isinstance(value, bool) or not isinstance(value, (str, int)) or not re.fullmatch(r"-?[0-9]+", str(value).strip()):
        raise ValueError(f"{field} 必须是整数")
    number = int(str(value).strip())
    if (minimum is not None and number < minimum) or (maximum is not None and number > maximum):
        raise ValueError(f"{field} 必须在 {minimum}–{maximum} 之间")
    return str(number)


def money(value, field, maximum):
    if isinstance(value, bool) or not isinstance(value, (str, int, Decimal)):
        raise ValueError(f"{field} 必须是金额字符串或 JSON 数值")
    raw = format(value, "f") if isinstance(value, Decimal) else str(value).strip()
    if not re.fullmatch(r"[0-9]+(?:\.[0-9]{1,2})?", raw):
        raise ValueError(f"{field} 必须是整数或最多两位小数，不可静默舍入")
    amount = Decimal(raw)
    if not Decimal("0") <= amount <= Decimal(str(maximum)):
        raise ValueError(f"{field} 必须在 0–{maximum} 元之间")
    return format(amount, ".2f")


def list_text(value, field, choices=None):
    parts = value.split(",") if isinstance(value, str) else value
    if not isinstance(parts, list) or not parts:
        raise ValueError(f"{field} 必须是非空文本或文本数组")
    parts = [text(part, field).strip() for part in parts]
    if choices is not None and any(part not in choices for part in parts):
        raise ValueError(f"{field} 含模板未列出的品类")
    return ",".join(parts)


def date_text(value, field, with_date=True):
    value = text(value, field).strip()
    pattern = r"[0-9]{4}-[0-9]{2}-[0-9]{2} [0-9]{2}:[0-9]{2}:[0-9]{2}" if with_date else r"[0-9]{2}:[0-9]{2}:[0-9]{2}"
    fmt = "%Y-%m-%d %H:%M:%S" if with_date else "%H:%M:%S"
    try:
        if not re.fullmatch(pattern, value):
            raise ValueError()
        datetime.strptime(value, fmt)
    except ValueError:
        raise ValueError(f"{field} 日期/时间格式或值不合法，须使用 {fmt}") from None
    return value


def prepare_records(records, kind, defaults):
    if not isinstance(records, list) or not records:
        raise ValueError("输入必须是包含至少一条活动的 JSON 数组")
    prepared, errors, used_defaults = [], [], Counter()
    unverified = {"backend_import", "city_and_population_entities"}
    required = REQUIRED + (["reward_percent", "max_amount_of_order"] if kind == "feizhang" else [])
    for index, original in enumerate(records, 1):
        try:
            if not isinstance(original, dict):
                raise ValueError("每条活动必须是英文字段对象")
            unknown = set(original) - (set(FIELDS[kind]) - {"id"})
            if unknown:
                raise ValueError("字段不属于当前模板或由脚本生成: " + ", ".join(sorted(unknown)))
            missing = [field for field in required if field not in original or original[field] is None or (isinstance(original[field], str) and not original[field].strip())]
            if missing:
                raise ValueError("缺少用户必填字段: " + ", ".join(missing))
            row = dict(defaults)
            row.update(original)
            row["id"] = str(index)
            row["title"] = text(row["title"], "title")
            row["city_ids"] = list_text(row["city_ids"], "city_ids")
            row["stock"] = integer(row["stock"], "stock", 1, 100000)
            row["price"] = money(row["price"], "price", 5000)
            row["get_type"] = text(row["get_type"], "get_type")
            row["attention"] = text(row["attention"], "attention", allow_empty=True)
            if row["tag"] not in ("普通卡", "早鸟卡"):
                raise ValueError("tag 仅支持普通卡或早鸟卡")
            if row["card_rule"] not in RULES[kind]:
                raise ValueError("card_rule 必须使用对应模板的固定规则选项")
            limit = row["order_limit"]
            row["order_limit"] = "" if row["card_rule"] == RULES[kind][0] and limit in (None, "") else integer(limit, "order_limit")
            row["order_category"] = list_text(row["order_category"], "order_category", CATEGORIES)
            for field in ("join_dri_tag", "forbid_dri_tag"):
                row[field] = text("" if row[field] is None else row[field], field, allow_empty=True)
                if not row[field].strip():
                    row[field] = ""
            if row["forbid_dri_tag"] and not row["join_dri_tag"]:
                raise ValueError("填写 forbid_dri_tag 时必须填写 join_dri_tag")
            if row["join_dri_tag"] or row["forbid_dri_tag"]:
                unverified.add("population_expiry")
            if kind == "feizhang":
                row["reward_percent"] = integer(row["reward_percent"], "reward_percent")
                row["max_amount_of_order"] = money(row["max_amount_of_order"], "max_amount_of_order", 100)
            for prefix in ("sale", "effect"):
                for edge in ("start", "end"):
                    field = f"{prefix}_{edge}_time"
                    row[field] = date_text(row[field], field)
                if row[f"{prefix}_start_time"] > row[f"{prefix}_end_time"]:
                    raise ValueError(f"{prefix}_start_time 不能晚于 {prefix}_end_time")
            pair_count = 0
            for pair in range(1, 6):
                begin, end = f"begin_time_{pair}", f"end_time_{pair}"
                start_value, end_value = row.get(begin, ""), row.get(end, "")
                start_value = text("" if start_value is None else start_value, begin, allow_empty=True).strip()
                end_value = text("" if end_value is None else end_value, end, allow_empty=True).strip()
                if bool(start_value) != bool(end_value):
                    raise ValueError(f"{begin} 与 {end} 必须成对提供")
                row[begin] = date_text(start_value, begin, False) if start_value else ""
                row[end] = date_text(end_value, end, False) if end_value else ""
                if start_value:
                    pair_count += 1
                    if start_value > end_value:
                        unverified.add("overnight_period_semantics")
            if not pair_count:
                raise ValueError("begin_time_1/end_time_1 等至少需要一组完整奖励时段")
            used_defaults.update(field for field in defaults if field not in original)
            prepared.append(row)
        except ValueError as exc:
            errors.append(f"第 {index} 条: {exc}")
    if errors:
        raise ValueError("；".join(errors))
    return prepared, dict(used_defaults), sorted(unverified)


def sheet_signature(ws, last_row=None):
    last_row = ws.max_row if last_row is None else last_row
    cells = [(cell.coordinate, None if cell.value == "" else cell.value, cell.number_format, copy(cell.font), copy(cell.fill), copy(cell.border), copy(cell.alignment), copy(cell.protection)) for row in ws.iter_rows(max_row=last_row) for cell in row]
    rows = [(index, dim.height, dim.hidden) for index, dim in ws.row_dimensions.items() if index <= last_row]
    columns = [(key, dim.min, dim.max, dim.width, dim.hidden) for key, dim in ws.column_dimensions.items()]
    return cells, rows, columns, str(ws.merged_cells), ws.sheet_state, ws.freeze_panes


def write_verified_workbook(path, wb, kind, records):
    ws = wb["活动信息"]
    expected_header = sheet_signature(ws, 2)
    expected_notes = sheet_signature(wb["重要说明"])
    prototype = [copy(cell._style) for cell in ws[3]]
    prototype_height = ws.row_dimensions[3].height
    for row_index, record in enumerate(records, 3):
        ws.row_dimensions[row_index].height = prototype_height
        for column, field in enumerate(FIELDS[kind], 1):
            cell = ws.cell(row_index, column)
            cell._style = copy(prototype[column - 1])
            cell.value = record[field]
            cell.data_type = "s"
            cell.number_format = "@"
    with path.open("xb") as stream:
        wb.save(stream)
    actual = load_workbook(path, data_only=False)
    try:
        if actual.sheetnames != wb.sheetnames or sheet_signature(actual["活动信息"], 2) != expected_header or sheet_signature(actual["重要说明"]) != expected_notes:
            raise ValueError("输出前两行或重要说明的内容/格式读回不一致")
        saved = actual["活动信息"]
        if saved.max_row != len(records) + 2 or saved.max_column != len(FIELDS[kind]):
            raise ValueError("输出活动行数或列数读回不一致")
        for row_index, record in enumerate(records, 3):
            for column, field in enumerate(FIELDS[kind], 1):
                cell = saved.cell(row_index, column)
                if (cell.value if cell.value is not None else "") != record[field] or cell.number_format != "@" or (record[field] and cell.data_type != "s"):
                    raise ValueError(f"输出第 {row_index} 行 {field} 的文本值/格式读回不一致")
    finally:
        actual.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description="按宏鹄模板生成飞涨卡/免佣卡 XLSX（仅新建）")
    parser.add_argument("--type", choices=tuple(FIELDS), required=True)
    parser.add_argument("--input", type=Path, required=True, help="UTF-8 JSON 活动数组")
    parser.add_argument("--output", type=Path, required=True, help="父目录已存在的新 XLSX 路径")
    parser.add_argument("--dry-run", action="store_true", help="只读校验，不创建目录、文件或日志")
    args = parser.parse_args(argv)
    wb = None
    try:
        if args.input.suffix.lower() != ".json" or args.output.suffix.lower() != ".xlsx":
            raise ValueError("输入必须是 .json，输出必须是 .xlsx")
        if args.output.exists() or args.output.is_symlink():
            raise ValueError("输出目标已存在或为符号链接，请使用新的文件名")
        if not args.output.parent.is_dir():
            raise ValueError("输出父目录不存在，请先按任务目录约定准备目录")
        wb, defaults = load_assets(args.type)
        with args.input.open(encoding="utf-8-sig") as stream:
            records = json.load(stream, parse_float=Decimal)
        prepared, used_defaults, unverified = prepare_records(records, args.type, defaults)
        if not args.dry_run:
            write_verified_workbook(args.output, wb, args.type, prepared)
        print(json.dumps({"status": "preview" if args.dry_run else "generated", "type": args.type, "activity_count": len(prepared), "column_count": len(FIELDS[args.type]), "defaults_used": used_defaults, "output": str(args.output.absolute()), "readback": "not-run" if args.dry_run else "passed", "backend_import": "not-performed", "unverified_checks": unverified}))
        return 0
    except (OSError, UnicodeError, ValueError, KeyError, BadZipFile) as exc:
        print(json.dumps({"status": "error", "error": str(exc)}))
        return 1
    finally:
        if wb is not None:
            wb.close()


if __name__ == "__main__":
    raise SystemExit(main())
