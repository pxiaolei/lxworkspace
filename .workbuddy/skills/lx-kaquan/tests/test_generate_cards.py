"""Test actual files and the CLI with fictional activities."""

import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


SKILL = Path(__file__).resolve().parents[1]
SCRIPT = SKILL / "scripts" / "generate_cards.py"


class CardTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix="lx-kaquan-test-"))
        self.input = self.root / "input.json"
        self.output = self.root / "result.xlsx"

    def record(self, **values):
        record = {
            "title": "虚构测试活动",
            "city_ids": "北京市",
            "price": "10.12",
            "sale_start_time": "2026-10-20 00:00:00",
            "sale_end_time": "2026-10-20 23:59:59",
            "effect_start_time": "2026-10-20 00:00:00",
            "effect_end_time": "2026-10-21 23:59:59",
            "begin_time_1": "07:00:00",
            "end_time_1": "08:30:00",
        }
        record.update(values)
        return record

    def flying(self, **values):
        values.setdefault("reward_percent", 5)
        values.setdefault("max_amount_of_order", "1")
        return self.record(**values)

    def run_cli(self, records=None, kind="mianyong", dry_run=False, output=None, script=SCRIPT, env=None):
        if records is not None:
            self.input.write_text(json.dumps(records, ensure_ascii=False), encoding="utf-8")
        command = [sys.executable, "-B", str(script), "--type", kind, "--input", str(self.input), "--output", str(output or self.output)]
        if dry_run:
            command.append("--dry-run")
        return subprocess.run(command, cwd=self.root, capture_output=True, text=True, encoding="utf-8", env=env)

    def success(self, result):
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual(result.stderr, "")
        return json.loads(result.stdout)

    def rejected(self, records, kind="mianyong", field=None):
        result = self.run_cli(records, kind)
        self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
        summary = json.loads(result.stdout)
        self.assertEqual(summary["status"], "error")
        if field:
            self.assertIn(field, summary["error"])
        self.assertFalse(self.output.exists())
        return summary

    def read_rows(self, output=None):
        wb = load_workbook(output or self.output, data_only=False)
        try:
            ws = wb["活动信息"]
            fields = [cell.value for cell in ws[2]]
            return [dict(zip(fields, [cell.value or "" for cell in row])) for row in ws.iter_rows(min_row=3)]
        finally:
            wb.close()

    def test_mianyong_defaults_include_complete_template_attention(self):
        summary = self.success(self.run_cli([self.record()]))
        row = self.read_rows()[0]
        expected = {"id": "1", "tag": "普通卡", "get_type": "购买", "stock": "999", "card_rule": "每天前x单免佣", "order_limit": "99", "order_category": "普通单,特赚单,稳收单", "join_dri_tag": "", "forbid_dri_tag": "", "price": "10.12"}
        for key, value in expected.items():
            self.assertEqual(row[key], value, key)
        template = load_workbook(SKILL / "assets" / "mianyong_template.xlsx")
        try:
            self.assertEqual(row["attention"], template["活动信息"]["AA3"].value)
        finally:
            template.close()
        self.assertIn("【免佣卡可退款规则】", row["attention"])
        self.assertEqual(summary["activity_count"], 1)
        self.assertEqual(summary["readback"], "passed")

    def test_feizhang_defaults_use_its_own_rule_and_attention(self):
        summary = self.success(self.run_cli([self.flying()], "feizhang"))
        row = self.read_rows()[0]
        self.assertEqual(row["card_rule"], "每天前x单飞涨")
        self.assertEqual(row["order_limit"], "99")
        self.assertEqual(row["stock"], "999")
        self.assertEqual(row["reward_percent"], "5")
        self.assertEqual(row["max_amount_of_order"], "1.00")
        self.assertEqual(row["tag"], "普通卡")
        self.assertEqual(row["get_type"], "购买")
        self.assertEqual(row["order_category"], "普通单,特赚单,稳收单")
        self.assertEqual(row["join_dri_tag"], "")
        self.assertEqual(row["forbid_dri_tag"], "")
        template = load_workbook(SKILL / "assets" / "feizhang_template.xlsx")
        try:
            self.assertEqual(row["attention"], template["活动信息"]["AC3"].value)
        finally:
            template.close()
        self.assertIn("【流水飞涨卡与其他活动或权益的兼容性】", row["attention"])
        self.assertEqual(summary["column_count"], 29)

    def test_headers_notes_and_styles_are_preserved_for_both_types(self):
        for kind, record in (("mianyong", self.record()), ("feizhang", self.flying())):
            with self.subTest(kind=kind):
                target = self.root / (kind + ".xlsx")
                self.success(self.run_cli([record], kind, output=target))
                template = load_workbook(SKILL / "assets" / (kind + "_template.xlsx"))
                actual = load_workbook(target)
                try:
                    self.assertEqual(actual.sheetnames, ["活动信息", "重要说明"])
                    for name in actual.sheetnames:
                        original_ws, actual_ws = template[name], actual[name]
                        last_row = 2 if name == "活动信息" else original_ws.max_row
                        for row in original_ws.iter_rows(max_row=last_row):
                            for cell in row:
                                saved = actual_ws[cell.coordinate]
                                self.assertEqual(saved.value, cell.value or None, cell.coordinate)
                                for attr in ("font", "fill", "border", "alignment", "protection"):
                                    self.assertEqual(copy(getattr(saved, attr)), copy(getattr(cell, attr)), (name, cell.coordinate, attr))
                                self.assertEqual(saved.number_format, cell.number_format)
                        self.assertEqual(str(actual_ws.merged_cells), str(original_ws.merged_cells))
                        self.assertEqual({k: v.width for k, v in actual_ws.column_dimensions.items()}, {k: v.width for k, v in original_ws.column_dimensions.items()})
                        for index in range(1, last_row + 1):
                            self.assertEqual(actual_ws.row_dimensions[index].height, original_ws.row_dimensions[index].height)
                    ws = actual["活动信息"]
                    self.assertEqual(ws.max_row, 3)
                    for cell in ws[3]:
                        self.assertEqual(cell.number_format, "@")
                        if cell.value is not None:
                            self.assertEqual(cell.data_type, "s")
                    self.assertEqual(ws["I3"].value, "购买" if kind == "feizhang" else "999")
                    self.assertEqual(ws["J3"].value, "999" if kind == "feizhang" else "购买")
                finally:
                    template.close()
                    actual.close()

    def test_multiple_records_start_at_row_three_with_new_sequences(self):
        self.success(self.run_cli([self.record(title="虚构活动一"), self.record(title="虚构活动二")]))
        rows = self.read_rows()
        self.assertEqual([row["id"] for row in rows], ["1", "2"])
        self.assertEqual([row["title"] for row in rows], ["虚构活动一", "虚构活动二"])
        self.assertEqual(rows[0]["begin_time_2"], "")
        self.assertEqual(rows[0]["end_time_2"], "")

    def test_user_values_override_defaults_and_zero_is_preserved(self):
        record = self.record(price=0, tag="早鸟卡", stock=1, get_type="领取", card_rule="累计前x单免佣", order_limit=2, order_category=["快车单", "增量单"], attention="用户文案\n保留换行", city_ids=["北京市", "上海市"])
        self.success(self.run_cli([record]))
        row = self.read_rows()[0]
        expected = {"price": "0.00", "tag": "早鸟卡", "stock": "1", "get_type": "领取", "card_rule": "累计前x单免佣", "order_limit": "2", "order_category": "快车单,增量单", "attention": "用户文案\n保留换行", "city_ids": "北京市,上海市"}
        for field, value in expected.items():
            self.assertEqual(row[field], value, field)

    def test_required_fields_are_not_filled_from_the_template_example(self):
        required = ["title", "city_ids", "price", "sale_start_time", "sale_end_time", "effect_start_time", "effect_end_time"]
        for field in required:
            for value in (None, ""):
                with self.subTest(field=field, value=value):
                    self.rejected([self.record(**{field: value})], field=field)
            record = self.record()
            record.pop(field)
            self.rejected([record], field=field)

    def test_feizhang_additional_fields_must_be_user_provided(self):
        for field in ("reward_percent", "max_amount_of_order"):
            record = self.flying()
            record.pop(field)
            self.rejected([record], "feizhang", field)
        self.rejected([self.flying(reward_percent="5.5")], "feizhang", "reward_percent")

    def test_price_stock_and_reward_cap_ranges_and_precision(self):
        for field, bad_values in (("price", ("-0.01", "5000.01", "1.001", True)), ("stock", (0, 100001, "1.5", True)), ("max_amount_of_order", ("-1", "100.01", "1.001"))):
            for value in bad_values:
                with self.subTest(field=field, value=value):
                    self.rejected([self.flying(**{field: value})], "feizhang", field)

    def test_numeric_json_money_keeps_decimal_precision(self):
        self.success(self.run_cli([self.flying(price=0.29, max_amount_of_order=0.1)], "feizhang"))
        row = self.read_rows()[0]
        self.assertEqual(row["price"], "0.29")
        self.assertEqual(row["max_amount_of_order"], "0.10")

    def test_boundary_values_are_accepted(self):
        self.success(self.run_cli([self.flying(price="5000.00", stock=100000, max_amount_of_order="100.00")], "feizhang"))
        self.assertEqual(self.read_rows()[0]["price"], "5000.00")

    def test_card_rule_tag_and_order_category_enums(self):
        for field, value in (("card_rule", "每天前x单飞涨"), ("tag", "不支持的标签"), ("order_category", "不支持的品类")):
            self.rejected([self.record(**{field: value})], field=field)

    def test_limited_rule_requires_integer_order_limit(self):
        for value in (None, "", "99.5"):
            self.rejected([self.record(order_limit=value)], field="order_limit")

    def test_single_order_rule_allows_blank_order_limit(self):
        self.success(self.run_cli([self.record(card_rule="单单免佣", order_limit="")]))
        self.assertEqual(self.read_rows()[0]["order_limit"], "")

    def test_at_least_one_complete_time_pair_is_required(self):
        for values in ({"begin_time_1": "", "end_time_1": ""}, {"begin_time_1": "", "end_time_1": "08:00:00"}, {"begin_time_1": "24:00:00"}, {"begin_time_1": "7:00"}):
            self.rejected([self.record(**values)], field="time")

    def test_fifth_time_pair_works_without_inheriting_example_periods(self):
        self.success(self.run_cli([self.record(begin_time_1="", end_time_1="", begin_time_5="22:00:00", end_time_5="01:00:00")]))
        row = self.read_rows()[0]
        self.assertEqual(row["begin_time_1"], "")
        self.assertEqual(row["begin_time_2"], "")
        self.assertEqual(row["begin_time_5"], "22:00:00")
        self.assertEqual(row["end_time_5"], "01:00:00")

    def test_sixth_pair_is_not_silently_added(self):
        self.rejected([self.record(begin_time_6="10:00:00", end_time_6="11:00:00")], field="begin_time_6")

    def test_invalid_dates_and_reversed_ranges_are_rejected(self):
        for values in ({"sale_start_time": "2026-02-30 00:00:00"}, {"sale_start_time": "2026/10/20"}, {"sale_end_time": "2026-10-19 23:59:59"}, {"effect_end_time": "2026-10-19 23:59:59"}):
            self.rejected([self.record(**values)], field="time")

    def test_forbidden_population_requires_join_population(self):
        self.rejected([self.record(forbid_dri_tag="虚构屏蔽人群")], field="join_dri_tag")

    def test_population_names_are_preserved_and_expiry_is_not_claimed_verified(self):
        summary = self.success(self.run_cli([self.record(join_dri_tag="虚构参与人群", forbid_dri_tag="虚构屏蔽人群")]))
        self.assertEqual(self.read_rows()[0]["join_dri_tag"], "虚构参与人群")
        self.assertIn("population_expiry", summary["unverified_checks"])

    def test_one_bad_record_prevents_all_output_and_does_not_log_titles(self):
        summary = self.rejected([self.record(), self.record(price="5001")], field="price")
        self.assertIn("2", summary["error"])
        self.assertNotIn("虚构测试活动", summary["error"])

    def test_dry_run_creates_no_files(self):
        self.input.write_text(json.dumps([self.record()], ensure_ascii=False), encoding="utf-8")
        before = set(self.root.iterdir())
        summary = self.success(self.run_cli(dry_run=True))
        self.assertEqual(summary["status"], "preview")
        self.assertEqual(summary["activity_count"], 1)
        self.assertEqual(set(self.root.iterdir()), before)

    def test_existing_output_is_not_overwritten_in_any_mode(self):
        self.output.write_bytes(b"existing business file")
        for dry_run in (False, True):
            result = self.run_cli([self.record()], dry_run=dry_run)
            self.assertEqual(result.returncode, 1)
            self.assertEqual(json.loads(result.stdout)["status"], "error")
            self.assertEqual(self.output.read_bytes(), b"existing business file")

    def test_missing_parent_is_not_created(self):
        output = self.root / "missing" / "result.xlsx"
        for dry_run in (True, False):
            result = self.run_cli([self.record()], output=output, dry_run=dry_run)
            self.assertEqual(result.returncode, 1)
            self.assertEqual(json.loads(result.stdout)["status"], "error")
            self.assertFalse(output.parent.exists())

    def test_broken_output_symlink_is_rejected(self):
        target = self.root / "target.xlsx"
        try:
            self.output.symlink_to(target)
        except (OSError, NotImplementedError):
            self.skipTest("symlinks unavailable")
        result = self.run_cli([self.record()])
        self.assertEqual(result.returncode, 1)
        self.assertEqual(json.loads(result.stdout)["status"], "error")
        self.assertFalse(target.exists())

    def test_empty_unknown_and_malformed_inputs_fail_without_output(self):
        for records in ([], {}, ["bad"], [self.record(unexpected="x")], [self.record(reward_percent=5)]):
            self.rejected(records)
        self.input.write_text("[invalid", encoding="utf-8")
        result = self.run_cli()
        self.assertEqual(result.returncode, 1)
        self.assertEqual(json.loads(result.stdout)["status"], "error")
        self.assertFalse(self.output.exists())

    def test_explicit_null_does_not_silently_reuse_defaults(self):
        for field in ("tag", "stock", "get_type", "card_rule", "attention"):
            self.rejected([self.record(**{field: None})], field=field)

    def test_text_starting_with_equals_does_not_become_a_formula(self):
        self.success(self.run_cli([self.record(title="=SUM(1,2)", attention="=原样文案")]))
        actual = load_workbook(self.output, data_only=False)
        try:
            self.assertEqual(actual["活动信息"]["B3"].value, "=SUM(1,2)")
            self.assertEqual(actual["活动信息"]["B3"].data_type, "s")
            self.assertEqual(actual["活动信息"]["AA3"].data_type, "s")
        finally:
            actual.close()

    def test_whole_skill_runs_after_copy_to_directory_with_spaces(self):
        copied = self.root / "portable skill"
        shutil.copytree(SKILL, copied, ignore=shutil.ignore_patterns("__pycache__"))
        self.success(self.run_cli([self.record()], script=copied / "scripts" / "generate_cards.py"))
        self.assertEqual(self.read_rows()[0]["stock"], "999")

    def test_ascii_console_preserves_chinese_output_path(self):
        output = self.root / "测试免佣卡.xlsx"
        summary = self.success(self.run_cli([self.record()], output=output, env=dict(os.environ, PYTHONIOENCODING="ascii")))
        self.assertEqual(summary["output"], str(output))

    def test_wrong_output_suffix_is_rejected(self):
        output = self.root / "result.csv"
        result = self.run_cli([self.record()], output=output)
        self.assertEqual(result.returncode, 1)
        self.assertEqual(json.loads(result.stdout)["status"], "error")
        self.assertFalse(output.exists())
